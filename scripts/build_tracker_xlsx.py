"""Generate the free net-worth tracker spreadsheet shipped at /net-worth-tracker-excel.

Run manually when the categories or the copy change:

    python scripts/build_tracker_xlsx.py

openpyxl is a **build-time** dependency only — deliberately not in
requirements.txt or pyproject, because the app never reads or writes xlsx at
runtime. The generated file is committed under app/static/ and served as a
static asset.

The categories come from `app.main._CALC_ASSETS` / `_CALC_LIABILITIES`, the same
curated slugs the web calculator uses, so the spreadsheet, the calculator and the
product itself can't drift apart.
"""

from __future__ import annotations

import pathlib
import sys

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))
from app import main as app_main  # noqa: E402

OUT = pathlib.Path(__file__).resolve().parent.parent / "app" / "static" / "networth-tracker.xlsx"

PERIODS = 12                      # three years of quarters
INK = "1F3A5F"                    # the app's brand navy
COPPER = "B87333"
# Indian digit grouping, which Excel has no locale for — crore, then lakh, then
# plain. Without this a spreadsheet for an Indian audience shows 15,720,000.
INR_FMT = '[>=10000000]"₹"##\\,##\\,##\\,##0;[>=100000]"₹"##\\,##\\,##0;"₹"##,##0'

thin = Side(style="thin", color="D8DEE7")


def _label(slug: str, rows: list[dict]) -> str:
    return next(r["title"] for r in rows if r["slug"] == slug)


def build() -> pathlib.Path:
    assets = app_main._calc_rows(app_main._CALC_ASSETS)
    liabilities = app_main._calc_rows(app_main._CALC_LIABILITIES)

    wb = Workbook()
    ws = wb.active
    ws.title = "Net worth"
    last_col = 1 + PERIODS
    last_letter = get_column_letter(last_col)

    ws["A1"] = "Net worth tracker"
    ws["A1"].font = Font(size=16, bold=True, color=INK)
    ws["A2"] = ("Fill in what each category is worth today. Change the first date below and "
                "the rest follow automatically.")
    ws["A2"].font = Font(size=10, italic=True, color="667085")

    # --- Date header: one real date, then every quarter derived from it --------
    hdr = 4
    ws.cell(hdr, 1, "Category").font = Font(bold=True, color="FFFFFF")
    ws.cell(hdr, 1).fill = PatternFill("solid", fgColor=INK)
    ws.cell(hdr, 2, "=DATE(2026,3,31)")
    for c in range(3, last_col + 1):
        # Self-populating, so the file doesn't go stale the year after it's built.
        ws.cell(hdr, c, f"=EDATE({get_column_letter(c - 1)}{hdr},3)")
    for c in range(2, last_col + 1):
        cell = ws.cell(hdr, c)
        cell.number_format = "mmm yyyy"
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.alignment = Alignment(horizontal="center")

    def section(title: str, rows: list[dict], start: int) -> tuple[int, int]:
        """Write a heading, the rows, and a subtotal. Returns (first, total_row)."""
        ws.cell(start, 1, title).font = Font(bold=True, size=11, color=COPPER)
        first = start + 1
        for i, r in enumerate(rows):
            ws.cell(first + i, 1, r["title"]).alignment = Alignment(indent=1)
            ws.cell(first + i, 1).font = Font(size=10)
            for c in range(2, last_col + 1):
                cell = ws.cell(first + i, c)
                cell.number_format = INR_FMT
                cell.border = Border(bottom=thin)
        total = first + len(rows)
        ws.cell(total, 1, f"Total {title.lower()}").font = Font(bold=True)
        for c in range(2, last_col + 1):
            col = get_column_letter(c)
            cell = ws.cell(total, c, f"=SUM({col}{first}:{col}{total - 1})")
            cell.number_format = INR_FMT
            cell.font = Font(bold=True)
            cell.border = Border(top=thin, bottom=thin)
        return first, total

    _, asset_total = section("Assets", assets, hdr + 1)
    _, liab_total = section("Liabilities", liabilities, asset_total + 2)

    net = liab_total + 2
    ws.cell(net, 1, "NET WORTH").font = Font(bold=True, size=12, color=INK)
    for c in range(2, last_col + 1):
        col = get_column_letter(c)
        cell = ws.cell(net, c, f"={col}{asset_total}-{col}{liab_total}")
        cell.number_format = INR_FMT
        cell.font = Font(bold=True, size=12, color=INK)
        cell.fill = PatternFill("solid", fgColor="EEF2F7")

    ws.column_dimensions["A"].width = 30
    for c in range(2, last_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 15
    ws.freeze_panes = "B5"

    # --- The point of a tracker is the line, not the number -------------------
    chart = LineChart()
    chart.title = "Net worth over time"
    chart.height, chart.width = 8, 24
    chart.add_data(Reference(ws, min_col=1, min_row=net, max_col=last_col, max_row=net),
                   from_rows=True, titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=2, min_row=hdr, max_col=last_col, max_row=hdr))
    chart.y_axis.numFmt = INR_FMT
    ws.add_chart(chart, f"A{net + 3}")

    # --- Sheet 2: how to use it, and the mistakes that make it wrong ----------
    how = wb.create_sheet("How to use")
    lines: list[tuple[str, bool]] = [
        ("How to use this tracker", True),
        ("", False),
        ("1. Set the first date in cell B4. Every other column follows it, a quarter apart.", False),
        ("2. Fill in what each category is worth TODAY — not what you paid for it.", False),
        ("3. Leave anything that doesn't apply blank. Totals and net worth are formulas;", False),
        ("   you don't need to touch them.", False),
        ("4. Come back once a quarter and fill the next column. The chart draws itself.", False),
        ("", False),
        ("Four things that make the number wrong", True),
        ("", False),
        ("Subtracting the same loan twice — enter the flat at full market value under", False),
        ("Assets and the outstanding home loan under Liabilities. Not the flat 'net of", False),
        ("the loan' with the loan listed as well.", False),
        ("", False),
        ("Using what you paid instead of what it's worth. Purchase price is history.", False),
        ("", False),
        ("Counting the EMI as the debt. What you owe is the outstanding balance, not", False),
        ("the monthly instalment and not the amount you originally borrowed.", False),
        ("", False),
        ("Leaving out EPF, PPF and jewellery. Illiquid is not the same as not owned —", False),
        ("for most salaried people these are a large share of the total.", False),
        ("", False),
        ("Tired of updating it by hand?", True),
        ("networthyhq.com imports your NSDL CAS and prices holdings live. Free.", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        cell = how.cell(i, 1, text)
        cell.font = Font(bold=bold, size=13 if bold and i == 1 else 11,
                         color=INK if bold else "333333")
    how.column_dimensions["A"].width = 95

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"wrote {path} ({path.stat().st_size:,} bytes)")
